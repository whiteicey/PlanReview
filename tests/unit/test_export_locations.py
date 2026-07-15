from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from app.domain.enums import BlockType, Origin, Severity
from app.domain.schemas import Finding, SourceSpan
from app.reports.exporters import export_anonymous_package, export_excel, export_word
from app.review.pipeline import ReviewRun, format_span_location


def test_format_span_location_table_cell():
    span = SourceSpan(
        span_id="D:t:1:8:1", document_id="D", section_path=["附件A关键参数表"],
        block_type=BlockType.TABLE_CELL, table_index=1, row_index=8, column_index=1,
        text="20", text_hash="h",
    )
    assert format_span_location(span) == "附件A关键参数表 表格 第9行第2列"


def test_format_span_location_paragraph():
    span = SourceSpan(
        span_id="D:p:6", document_id="D", section_path=["开发部署方案"],
        block_type=BlockType.PARAGRAPH, paragraph_index=6, text="…", text_hash="h",
    )
    assert format_span_location(span) == "开发部署方案 第7段"


def _run_with_location() -> ReviewRun:
    return ReviewRun(
        case_id="CASE-loc",
        findings=[
            Finding(
                finding_id="F1", origin=Origin.RULE, category="aggregation",
                severity=Severity.HIGH, parameter="开发井总数", title="开发井总数需复核",
                description="求和不一致", suggestion="请补充证据并由专家复核",
                rule_id="CONSISTENCY-002", evidence_span_ids=["D:t:1:8:1", "D:p:6"],
                needs_human_review=True,
            )
        ],
        evidence_locations={"D:t:1:8:1": "附件A关键参数表 表格 第9行第2列", "D:p:6": "开发部署方案 第7段"},
    )


def test_excel_export_includes_location_column():
    run = _run_with_location()
    target = __import__("pathlib").Path("storage") / "test_loc.xlsx"
    export_excel(run, target)
    workbook = load_workbook(target)
    sheet = workbook.active
    header = [cell.value for cell in sheet[2]]
    assert "location" in header
    location_col = header.index("location")
    value = sheet[3][location_col].value
    assert "附件A关键参数表 表格 第9行第2列" in value
    assert "开发部署方案 第7段" in value
    target.unlink(missing_ok=True)


def test_word_export_includes_location():
    run = _run_with_location()
    target = __import__("pathlib").Path("storage") / "test_loc.docx"
    export_word(run, target)
    from docx import Document

    text = "\n".join(p.text for p in Document(target).paragraphs)
    assert "问题位置：" in text
    assert "附件A关键参数表 表格 第9行第2列" in text
    target.unlink(missing_ok=True)


def test_anonymous_package_omits_readable_location():
    run = _run_with_location()
    target = __import__("pathlib").Path("storage") / "test_loc_anon.zip"
    export_anonymous_package(run, target)
    with ZipFile(target) as archive:
        blob = archive.read("anonymous-findings.json").decode("utf-8")
    # Human-readable locations must not leak into the de-identified package.
    assert "附件A关键参数表" not in blob
    assert "开发部署方案" not in blob
    target.unlink(missing_ok=True)
